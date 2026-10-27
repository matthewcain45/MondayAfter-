#!/usr/bin/env python3
"""All the words in The Monday After, in one spreadsheet.

  python3 scripts/app-text.py export [docs/app-text.xlsx]
      Writes every piece of reader-facing text to a spreadsheet, one row each.

  python3 scripts/app-text.py apply docs/app-text.xlsx
      Puts the "New wording" column back into the source files. Rows left
      blank are skipped. A row is refused if its "Current wording" no longer
      matches the source (the spreadsheet is out of date: export a fresh one).

Uses openpyxl (a well-known spreadsheet library). It only runs in a Claude
session or on a developer's machine, never on the website.
"""
import json
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
SCENARIOS = ["site/scenarios/incident-file-01.json"]

# Scenario fields that are settings or rules, not words a reader sees.
NOT_TEXT = {"id", "type", "keyStyle", "tone", "effect", "picked", "pulled", "choice", "thread",
            "defaultWhen", "detailFrom", "key", "replyAddress", "progress"}

SCREEN_NAMES = {
    "scene": "Screen 1 · The scene", "d1": "Screen 2 · Decision 1", "holding": "Screen 3 · Holding position",
    "reveal1": "Screen 4 · Reveal", "twist": "Screen 5 · Saturday, the twist", "d2": "Screen 6 · Decision 2",
    "recorder": "Screen 8 · Flight recorder", "forecast": "Screen 9 · Forecast", "ledger": "Screen 10 · Ledger",
    "takeback": "Screen 11 · Take one thing back", "closing": "Screen 12 · Closing",
}

PAGES = [
    ("site/signup.html", "Sign-up page"),
    ("site/confirm.html", "Confirm page"),
    ("site/manage.html", "Manage page (pause, unsubscribe, delete)"),
    ("site/index.html", "Scenario app (loading)"),
    ("site/404.html", "Page not found"),
    ("site/privacy.html", "Privacy notice"),
]

# Words written into the code. (file, text, where it shows)
CODE = [
    ("site/js/app.js", "Open this file from your email.", "App · opened without a personal link"),
    ("site/js/app.js", "Each Monday After file opens from the personal link in your email.", "App · opened without a personal link"),
    ("site/js/app.js", "Not subscribed yet? Sign up", "App · opened without a personal link"),
    ("site/js/app.js", "This link has expired.", "App · expired link"),
    ("site/js/app.js", "Reply to any Monday After email and we'll send you a fresh one.", "App · expired link"),
    ("site/js/app.js", "This file couldn't be opened.", "App · couldn't open"),
    ("site/js/app.js", "Check your connection and try the link in your email again.", "App · couldn't open"),
    ("site/js/app.js", "Please try the link in your email again.", "App · couldn't open"),
    ("site/js/app.js", "Something went wrong.", "App · general error"),
    ("site/js/app.js", "Not saved yet. Check your connection; we'll try again.", "App · saving problem"),
    ("site/js/app.js", "You're assuming…", "App · reveal box heading"),
    ("site/js/app.js", "Gain.", "App · thread answer"),
    ("site/js/app.js", "Cost.", "App · thread answer"),
    ("site/js/app.js", "Lock in", "App · thread button"),
    ("site/js/app.js", "Back to the file", "App · thread button"),
    ("site/js/app.js", "Back to the ledger", "App · ledger detail and take-back buttons"),
    ("site/js/app.js", "The Monday After · Thread", "App · ledger detail header"),
    ("site/js/app.js", "Three months on", "App · ledger detail"),
    ("site/js/app.js", "Next time", "App · ledger detail"),
    ("site/js/app.js", "Assured on paper", "App · ledger detail"),
    ("site/js/app.js", "You can pick up to ", "App · too many flags (then a number)"),
    ("site/js/app.js", ". Untick one first.", "App · too many flags (end)"),
    ("site/js/app.js", "You saved your flight recorder on ", "App · flight recorder (then a date)"),
    ("site/js/app.js", ". Saving again replaces it.", "App · flight recorder (end)"),
    ("site/js/app.js", "Please write something for the first hour, or skip.", "App · flight recorder"),
    ("site/js/app.js", "Couldn't save. Check your connection and try again.", "App · flight recorder"),
    ("site/js/app.js", "Pick one thing to take back, or skip.", "App · take one thing back"),
    ("site/js/app.js", "Sent. Check your inbox.", "App · email me my flight recorder"),
    ("site/js/app.js", "Couldn't send. Check your connection and try again.", "App · email me my flight recorder"),
    ("site/js/app.js", ": I'd check this in my own organisation", "App · read aloud by screen readers on the flag button"),
    ("site/js/app.js", ": open three months on and next time", "App · read aloud by screen readers on ledger rows"),
    ("site/js/pages.js", "Something went wrong. Please try again.", "Sign-up, confirm and manage · general error"),
    ("site/js/pages.js", "Couldn't reach us. Check your connection and try again.", "Sign-up and manage · no connection"),
    ("site/js/pages.js", "Couldn't reach us. Check your connection and reload this page.", "Confirm and manage · no connection"),
    ("site/js/pages.js", "This confirm link has expired or was already used. Sign up again and we'll send a fresh one.", "Confirm page · expired link"),
    ("site/js/pages.js", "You're unsubscribed, so no more files will arrive. You can still delete your data below.", "Manage page · unsubscribed"),
    ("site/js/pages.js", "You're on a break from the monthly files.", "Manage page · on a break"),
    ("site/js/pages.js", "You get one incident file a month.", "Manage page · subscribed"),
    ("site/js/pages.js", "No files until ", "Manage page · on a break (then a date)"),
    ("site/js/pages.js", ". They start again by themselves after that.", "Manage page · on a break (end)"),
    ("site/js/pages.js", "This link has expired or isn't valid. Use the link at the bottom of your most recent Monday After email.", "Manage page · expired link"),
    ("backend/src/emails.ts", "The Monday After, from Capstone Global Advisory.", "Every email · footer"),
    ("backend/src/emails.ts", "Take a break, unsubscribe or delete your data: ", "Every email · footer (then the link)"),
    ("backend/src/emails.ts", "Confirm your email for ", "Confirm email · subject (then the name)"),
    ("backend/src/emails.ts", "Someone (hopefully you) asked to receive ", "Confirm email (then the name)"),
    ("backend/src/emails.ts", ", a short monthly incident file for technology leaders.", "Confirm email"),
    ("backend/src/emails.ts", "Confirm your email here:", "Confirm email"),
    ("backend/src/emails.ts", "This link works for 24 hours. If you didn't ask for this, ignore this email and nothing more will happen: we delete unconfirmed sign-ups after 7 days.", "Confirm email"),
    ("backend/src/emails.ts", "This link is personal to you. Please don't forward it; use Share this file at the end instead.", "Monthly file email · under the link"),
    ("backend/src/emails.ts", "Your flight recorder: ", "Flight recorder email · subject (then the file name)"),
    ("backend/src/emails.ts", "Here is what you wrote in your flight recorder for ", "Flight recorder email (then the file name and date)"),
    ("backend/src/emails.ts", "When the next incident hits, compare it with what you actually do.", "Flight recorder email"),
    ("backend/src/emails.ts", "(left blank)", "Flight recorder email · an empty answer"),
    ("backend/src/emails.ts", "Only you can see this. The author of ", "Flight recorder email (then the name)"),
    ("backend/src/emails.ts", " can't read your flight recorder.", "Flight recorder email (end)"),
    ("backend/src/service.ts", "This link has expired or isn't valid. Reply to any Monday After email and we'll send you a fresh one.", "Message · expired link"),
    ("backend/src/service.ts", "Too many requests. Please wait a little and try again.", "Message · too many tries"),
    ("backend/src/service.ts", "We're getting a lot of sign-ups right now. Please try again in an hour.", "Sign-up page · busy"),
    ("backend/src/service.ts", "There's no saved flight recorder for this file.", "App · email me my flight recorder, nothing saved"),
    ("backend/src/service.ts", "You're unsubscribed, so there's nothing to pause.", "Manage page"),
    ("backend/src/validate.ts", "Email is missing", "Sign-up page · empty email"),
    ("backend/src/validate.ts", "That doesn't look like an email address", "Sign-up page · bad email"),
    ("backend/src/validate.ts", "Please keep each entry under ", "Flight recorder · too long (then a number)"),
    ("backend/src/validate.ts", "Please write something for the first hour", "Flight recorder · first hour empty"),
]

TABS = ["Incident File 01", "App messages", "Pages", "Emails", "Privacy notice"]
HEAD = ["Row", "Where it shows", "Which part", "Current wording", "New wording", "Notes"]
HIDDEN = ["file", "locator"]


def read(rel):
    return (ROOT / rel).read_text(encoding="utf-8")


def write(rel, text):
    (ROOT / rel).write_text(text, encoding="utf-8")


# ---------- scenario files ----------

def scenario_rows(rel):
    data = json.loads(read(rel))
    rows = []

    def where(path):
        if path[0] == "email":
            return "Monthly email"
        if path[0] == "screens":
            s = data["screens"][path[1]]
            return SCREEN_NAMES.get(s.get("id"), "Screen · " + str(s.get("id")))
        if path[0] == "threads":
            key = path[1]
            th = data["threads"][key] if isinstance(data["threads"], dict) else data["threads"][key]
            title = th.get("title") if isinstance(th.get("title"), str) else th.get("id", key)
            return "Thread · " + str(title)
        return "File details"

    def part(path):
        bits = []
        rest = path[2:] if path[0] in ("screens", "threads") else path[1:]
        for p in rest:
            if isinstance(p, int):
                bits.append("#" + str(p + 1))
            else:
                bits.append(p)
        return " › ".join(bits)

    def walk(o, path):
        if isinstance(o, dict):
            if set(o) == {"todo"}:
                rows.append([where(path), part(path), "", "COPY NEEDED: " + o["todo"], json.dumps(path)])
                return
            for k, v in o.items():
                if k in NOT_TEXT:
                    continue
                walk(v, path + [k])
        elif isinstance(o, list):
            for i, v in enumerate(o):
                walk(v, path + [i])
        elif isinstance(o, str) and o.strip():
            rows.append([where(path), part(path), o, "", json.dumps(path)])

    walk(data, [])
    return [(rel, *r) for r in rows]


def set_path(data, path, value):
    for p in path[:-1]:
        data = data[p]
    data[path[-1]] = value


def get_path(data, path):
    for p in path:
        data = data[p]
    return data


# ---------- web pages ----------

INLINE = {"b", "strong", "i", "em", "a", "span", "br", "small", "abbr", "code"}
VOID = {"meta", "link", "input", "br", "img", "hr", "source"}
SKIP = {"script", "style"}
ATTRS = {"placeholder", "aria-label", "alt", "title"}


class Node:
    def __init__(self, tag, attrs, start, raw):
        self.tag, self.attrs, self.start, self.raw = tag, attrs, start, raw
        self.inner_start = start + len(raw)
        self.inner_end = None
        self.children = []


class Tree(HTMLParser):
    def __init__(self, src):
        super().__init__(convert_charrefs=False)
        self.src = src
        self.lines = [0]
        for m in re.finditer("\n", src):
            self.lines.append(m.end())
        self.root = Node("root", [], 0, "")
        self.stack = [self.root]
        self.feed(src)

    def pos(self):
        line, col = self.getpos()
        return self.lines[line - 1] + col

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs, self.pos(), self.get_starttag_text())
        self.stack[-1].children.append(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].children.append(Node(tag, attrs, self.pos(), self.get_starttag_text()))

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                self.stack[i].inner_end = self.pos()
                del self.stack[i:]
                return

    def handle_data(self, data):
        self.stack[-1].children.append((self.pos(), data))


def all_inline(node):
    for c in node.children:
        if isinstance(c, Node):
            if c.tag not in INLINE or not all_inline(c):
                return False
    return True


def page_rows(rel, label):
    src = read(rel)
    tree = Tree(src)
    found = []  # (offset, which part, display, locator)

    def text_row(start, end, what):
        inner = src[start:end]
        if not re.search(r"[A-Za-z]", inner):
            return
        found.append((start, what, re.sub(r"\s+", " ", inner).strip(), {"kind": "text", "raw": ">" + inner + "<"}))

    def attrs_of(node):
        names = {"placeholder": "box hint", "aria-label": "read aloud by screen readers", "alt": "image description", "title": "tooltip"}
        for k, v in node.attrs:
            if k in ATTRS and v and re.search(r"[A-Za-z]", v):
                found.append((node.start, names[k], v, {"kind": "attr", "raw": f'{k}="{v}"'}))
        d = dict(node.attrs)
        if node.tag == "meta" and d.get("name") == "description":
            found.append((node.start, "search result summary", d["content"], {"kind": "attr", "raw": f'content="{d["content"]}"'}))

    def visit(node):
        if node.tag in SKIP:
            return
        attrs_of(node)
        if node.tag == "title":
            text_row(node.inner_start, node.inner_end, "browser tab title")
            return
        if dict(node.attrs).get("aria-hidden") == "true":
            return
        has_text = any(not isinstance(c, Node) and re.search(r"[A-Za-z]", c[1]) for c in node.children)
        if node.tag != "root" and node.inner_end is not None and (has_text or node.tag in INLINE) and all_inline(node):
            text_row(node.inner_start, node.inner_end, node.tag)
            for c in node.children:
                if isinstance(c, Node):
                    for sub in walk_attrs(c):
                        attrs_of(sub)
            return
        for c in node.children:
            if isinstance(c, Node):
                visit(c)
            else:
                start, data = c
                text_row(start, start + len(data), "text")

    def walk_attrs(n):
        yield n
        for c in n.children:
            if isinstance(c, Node):
                yield from walk_attrs(c)

    visit(tree.root)
    found.sort(key=lambda f: f[0])
    names = {"h1": "heading", "h2": "heading", "h3": "heading", "p": "paragraph", "li": "list item",
             "button": "button", "label": "label", "a": "link", "span": "text", "div": "text",
             "figcaption": "caption", "footer": "footer"}
    rows, seen = [], set()
    for _, what, display, loc in found:
        if loc["raw"] in seen:
            continue
        seen.add(loc["raw"])
        rows.append((rel, label, names.get(what, what), display, "", json.dumps(loc)))
    return rows


# ---------- words in the code ----------

def code_rows():
    rows = []
    for rel, text, where in CODE:
        raw = json.dumps(text, ensure_ascii=False)
        if raw not in read(rel):
            sys.exit(f"Can't find {raw} in {rel}. Update the CODE list in scripts/app-text.py.")
        rows.append((rel, where, "message", text, "", json.dumps({"kind": "code", "raw": raw})))
    return rows


# ---------- export ----------

def export(out):
    tabs = {name: [] for name in TABS}
    for rel in SCENARIOS:
        tabs["Incident File 01"] += scenario_rows(rel)
    tabs["App messages"] += [r for r in code_rows() if not r[0].endswith("emails.ts")]
    tabs["Emails"] += [r for r in code_rows() if r[0].endswith("emails.ts")]
    for rel, label in PAGES:
        tabs["Privacy notice" if "privacy" in rel else "Pages"] += page_rows(rel, label)

    wb = Workbook()
    guide = wb.active
    guide.title = "How to use"
    for line in [
        "All the words in The Monday After",
        "",
        "1. Each tab holds one area: the scenario, the app's messages, the web pages, the emails and the privacy notice.",
        "2. To change some wording, type your new version in the yellow 'New wording' column. Leave the rest blank.",
        "3. Keep anything in <angle brackets> (bold text and links) and anything in {curly brackets} (it's filled in automatically).",
        "4. Rows marked COPY NEEDED haven't been written yet: type the words in 'New wording'.",
        "5. Don't change the 'Current wording' column, and don't add, delete or reorder rows.",
        "6. Send the file back to Claude. Claude puts your changes in, checks them, and you press Deploy test copy.",
        "",
        "Some sentences are split into pieces because the app adds a number, a date or a name in between.",
        "The 'Where it shows' column says which piece is which.",
        "Changing the privacy notice? Remember to update the 'Last updated' date, and tell Claude if what we store or how long we keep it has changed.",
    ]:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 120

    yellow = PatternFill("solid", fgColor="FFF4C2")
    n = 0
    for name in TABS:
        ws = wb.create_sheet(name)
        ws.append(HEAD + HIDDEN)
        for c in ws[1]:
            c.font = Font(bold=True)
        for rel, where, what, current, notes, loc in tabs[name]:
            n += 1
            ws.append([n, where, what, current, "", notes, rel, loc])
        for col, width in zip("ABCDEF", [6, 30, 24, 70, 70, 30]):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for c in row[:6]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            row[4].fill = yellow
        ws.column_dimensions["G"].hidden = True
        ws.column_dimensions["H"].hidden = True
        ws.freeze_panes = "B2"
    wb.save(out)
    print(f"Wrote {n} rows to {out}")


# ---------- apply ----------

def apply(path):
    wb = load_workbook(path)
    changes = []
    for name in TABS:
        if name not in wb.sheetnames:
            continue
        for row in wb[name].iter_rows(min_row=2, values_only=True):
            num, where, what, current, new, notes, rel, loc = (list(row) + [None] * 8)[:8]
            if new is None or str(new).strip() == "":
                continue
            changes.append((num, rel, json.loads(loc), current or "", str(new).strip()))

    problems, done, renamed = [], 0, []
    scen = {}
    texts = {}
    for num, rel, loc, current, new in changes:
        # Keep the spaces at either end of a sentence piece, which the app joins to a name or date.
        new = current[: len(current) - len(current.lstrip())] + new + current[len(current.rstrip()):]
        if rel in SCENARIOS:
            data = scen.setdefault(rel, json.loads(read(rel)))
            now = get_path(data, loc)
            now_text = "" if isinstance(now, dict) else now
            if now_text != current:
                problems.append(f"Row {num}: the wording has changed since this spreadsheet was made.")
                continue
            set_path(data, loc, new)
            done += 1
            continue
        src = texts.setdefault(rel, read(rel))
        raw = loc["raw"]
        if raw not in src:
            problems.append(f"Row {num}: the wording has changed since this spreadsheet was made.")
            continue
        if loc["kind"] == "code":
            rep = json.dumps(new, ensure_ascii=False)
            renamed.append((rel, raw, rep))
        elif loc["kind"] == "attr":
            name = raw.split("=", 1)[0]
            rep = f'{name}="' + new.replace("&", "&amp;").replace('"', "&quot;") + '"'
        else:
            inner = raw[1:-1]
            lead = inner[: len(inner) - len(inner.lstrip())]
            trail = inner[len(inner.rstrip()):]
            if inner.count("<") != new.count("<"):
                problems.append(f"Row {num}: the <angle bracket> parts don't match the current wording, so it was skipped.")
                continue
            rep = ">" + lead + new + trail + "<"
        texts[rel] = src.replace(raw, rep)
        done += 1

    for rel, data in scen.items():
        write(rel, json.dumps(data, indent=2, ensure_ascii=False) + "\n")
    for rel, src in texts.items():
        write(rel, src)
    # Keep this script's own list of code messages in step with the new wording.
    me = Path(__file__).read_text(encoding="utf-8")
    for rel, raw, rep in renamed:
        me = me.replace(f'("{rel}", {raw},', f'("{rel}", {rep},')
    Path(__file__).write_text(me, encoding="utf-8")
    print(f"Applied {done} change(s).")
    for p in problems:
        print("  " + p)


if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in ("export", "apply"):
        sys.exit(__doc__)
    if sys.argv[1] == "export":
        export(sys.argv[2] if len(sys.argv) > 2 else str(ROOT / "docs/app-text.xlsx"))
    else:
        apply(sys.argv[2])
